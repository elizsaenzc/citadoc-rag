import io

import docx
import pytest

from app.ingestion.chunker import chunk_text
from app.ingestion.loaders import UnsupportedFileError, load_document


def test_short_text_is_one_chunk():
    assert chunk_text("Hola mundo.", size=100, overlap=10) == ["Hola mundo."]


def test_empty_text_has_no_chunks():
    assert chunk_text("   \n\n  ", size=100, overlap=10) == []


def test_chunks_respect_size_and_overlap():
    paragraphs = [f"Párrafo número {i} con algo de contenido para rellenar." for i in range(40)]
    chunks = chunk_text("\n\n".join(paragraphs), size=300, overlap=60)
    assert len(chunks) > 1
    assert all(len(c) <= 300 for c in chunks)
    # consecutive chunks share text thanks to the overlap
    assert chunks[1][:15] in chunks[0]


def test_very_long_sentence_is_hard_split():
    chunks = chunk_text("a" * 2500, size=1000, overlap=100)
    assert all(len(c) <= 1000 for c in chunks)
    assert "".join(chunks).count("a") >= 2500


def test_invalid_overlap():
    with pytest.raises(ValueError):
        chunk_text("texto", size=100, overlap=100)


def test_load_txt_and_md():
    pages = load_document("notas.md", "# Título\n\nContenido".encode())
    assert pages[0].number is None
    assert "Contenido" in pages[0].text


def test_load_docx_includes_tables():
    document = docx.Document()
    document.add_paragraph("Política de gastos")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Dietas"
    table.cell(0, 1).text = "30 euros"
    buffer = io.BytesIO()
    document.save(buffer)
    text = load_document("gastos.docx", buffer.getvalue())[0].text
    assert "Política de gastos" in text
    assert "Dietas | 30 euros" in text


def test_unsupported_extension():
    with pytest.raises(UnsupportedFileError):
        load_document("hoja.xlsx", b"data")


def _make_pdf(pages: list[str]) -> bytes:
    """Build a minimal valid PDF with one line of text per page (no extra deps)."""
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        None,  # the page tree, filled in once the pages exist
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    kids = []
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        content_id = len(objects)
        objects.append(
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content_id} 0 R /Resources << /Font << /F1 3 0 R >> >> >>"
        )
        kids.append(f"{len(objects)} 0 R")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out, offsets = b"%PDF-1.4\n", []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def test_load_pdf_keeps_page_numbers():
    pages = load_document("informe.pdf", _make_pdf(["Primera pagina", "Segunda pagina"]))
    assert [p.number for p in pages] == [1, 2]
    assert "Segunda" in pages[1].text
