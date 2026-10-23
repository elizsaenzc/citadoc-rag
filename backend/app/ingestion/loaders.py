"""Extract text from uploaded files, keeping page numbers when the format has them."""

from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class UnsupportedFileError(ValueError):
    pass


@dataclass
class Page:
    number: int | None  # None when the format has no pages (txt, md, docx)
    text: str


def load_document(filename: str, content: bytes) -> list[Page]:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileError(
            f"Formato no soportado: {ext or 'sin extensión'}. Usa PDF, DOCX, TXT o MD."
        )
    if ext == ".pdf":
        return _load_pdf(content)
    if ext == ".docx":
        return _load_docx(content)
    return [Page(number=None, text=_decode_text(content))]


def _load_pdf(content: bytes) -> list[Page]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(Page(number=i, text=text))
    return pages


def _load_docx(content: bytes) -> list[Page]:
    import docx

    document = docx.Document(io.BytesIO(content))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return [Page(number=None, text="\n\n".join(parts))]


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")
